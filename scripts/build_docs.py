"""Génère docs/test-cases.xlsx et docs/bug-reports/*.md à partir de testcases.py, bugs.py et results.json."""
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from bugs import BUGS, ENV
from testcases import CASES

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
RES = json.loads((ROOT / "scripts" / "results.json").read_text(encoding="utf-8"))
BUG_OF = {tc: b["id"] for b in BUGS for tc in b["tc"]}

STATUS_FILL = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFEB9C", "Skipped": "D9D9D9"}
HEAD = PatternFill("solid", fgColor="1F2937")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def sheet(ws, headers, rows, widths):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.border = BORDER
            if c.row > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def build_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "Cas de test"
    headers = ["ID", "Type", "Canal", "Titre", "Objectif", "Préconditions", "Données", "Étapes",
               "Résultat attendu", "Résultat obtenu", "Statut", "Preuves", "Bug lié", "Risque"]
    rows = []
    for c in CASES:
        r = RES[c["id"]]
        rows.append([c["id"], c["type"], c["canal"], c["titre"], c["objectif"], c["preconditions"], c["donnees"],
                     c["etapes"], c["attendu"], r["obtenu"], r["statut"], "\n".join(r["preuves"]),
                     BUG_OF.get(c["id"], ""), c["risque"]])
    sheet(ws, headers, rows, [9, 12, 12, 34, 32, 28, 30, 42, 38, 48, 9, 36, 9, 7])
    for row in ws.iter_rows(min_row=2):
        st = row[10]
        st.fill = PatternFill("solid", fgColor=STATUS_FILL[st.value])
        st.font = Font(bold=True)

    # Synthèse
    s = wb.create_sheet("Synthèse")
    statuses = ["Pass", "Fail", "Blocked", "Skipped"]
    types = ["Nominal", "Limite", "Erreur", "Règle métier"]
    cnt = Counter((c["type"], RES[c["id"]]["statut"]) for c in CASES)
    srows = [[t] + [cnt[(t, st)] for st in statuses] + [sum(cnt[(t, st)] for st in statuses)] for t in types]
    srows.append(["TOTAL"] + [sum(cnt[(t, st)] for t in types) for st in statuses] + [len(CASES)])
    total = len(CASES)
    srows.append(["% du total"] + [f"{100 * srows[-1][i + 1] / total:.0f} %" for i in range(4)] + ["100 %"])
    sheet(s, ["Type", *statuses, "Total"], srows, [16, 10, 10, 10, 10, 10])
    for c in s[len(srows) + 1 - 1] + s[len(srows) + 1]:
        c.font = Font(bold=True)

    # Bugs
    b = wb.create_sheet("Bugs")
    sheet(b, ["ID", "Titre", "Zone", "Sévérité", "Priorité", "Cas de test", "Rapport"],
          [[x["id"], x["titre"], x["zone"], x["sev"], x["prio"], ", ".join(x["tc"]), f"bug-reports/{x['id']}.md"] for x in BUGS],
          [9, 70, 24, 11, 9, 18, 26])
    wb.save(DOCS / "test-cases.xlsx")


def build_bugs():
    out = DOCS / "bug-reports"
    out.mkdir(parents=True, exist_ok=True)
    index = ["# Bug reports — campagne initiale QuickOrder V3.1", "",
             f"{len(BUGS)} anomalies reproduites le 13/10/2026. Échelles : sévérité Bloquante > Critique > Majeure > Mineure ; "
             "priorité P1 (à corriger avant mise en production), P2 (avant MEP ou correctif immédiat), P3 (backlog).", "",
             "| ID | Titre | Sévérité | Priorité | Cas de test |", "|---|---|---|---|---|"]
    for x in BUGS:
        index.append(f"| [{x['id']}]({x['id']}.md) | {x['titre']} | {x['sev']} | {x['prio']} | {', '.join(x['tc'])} |")
        steps = "\n".join(f"{i}. {s}" for i, s in enumerate(x["etapes"], 1))
        proofs = "\n".join(
            (f"![{p}](../evidence/{p})" if p.endswith(".png") else f"- [`{p}`](../evidence/{p})") for p in x["preuves"])
        md = f"""# {x['id']} — {x['titre']}

| Champ | Valeur |
|---|---|
| **Sévérité** | {x['sev']} |
| **Priorité** | {x['prio']} |
| **Zone** | {x['zone']} |
| **Cas de test** | {', '.join(x['tc'])} |
| **Statut** | Ouvert — reproduit à chaque exécution de la campagne (base réinitialisée) |
| **Environnement** | {ENV} |

## Contexte
{x['contexte']}

## Préconditions
{x['precond']}

## Étapes de reproduction
{steps}

## Résultat attendu
{x['attendu']}

## Résultat obtenu
{x['obtenu']}

## Preuves
{proofs}

## Analyse (cause probable dans le code)
{x['cause']}
"""
        (out / f"{x['id']}.md").write_text(md, encoding="utf-8")
    (out / "README.md").write_text("\n".join(index) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_xlsx()
    build_bugs()
    c = Counter(r["statut"] for r in RES.values())
    print(len(CASES), "cas", dict(c), "|", len(BUGS), "bugs")
